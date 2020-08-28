import numpy as np
import pandas as pd
import xlrd
import docx
import re
from openpyxl import load_workbook
from operator import itemgetter
mydoc=docx.Document()

def toDictionary(filename,sheetname,custom_row=False,custom_col=False,**kwargs):
    flag=0
    if(custom_row&custom_col):
        df = pd.read_excel(filename, sheet_name=sheetname,header=None).iloc[kwargs['row_start']:kwargs['row_end'],kwargs['col_start']:kwargs['col_end']]
        flag=1
        col_start=kwargs['col_start']
        col_end=kwargs['col_end']
        row_start=kwargs['row_start']
        row_end=kwargs['row_end']

    elif(custom_row):
        df = pd.read_excel(filename, sheet_name=sheetname,header=None).iloc[kwargs['row_start']:kwargs['row_end'],:]
        flag=2
        col_start=0
        col_end=df.shape[1]
        row_start=kwargs['row_start']
        row_end=kwargs['row_end']
        
    elif(custom_col):
        df = pd.read_excel(filename, sheet_name=sheetname,header=None).iloc[:,kwargs['col_start']:kwargs['col_end']]
        flag=3
        row_start=0
        row_end=df.shape[0]
        col_start=kwargs['col_start']
        col_end=kwargs['col_end']
    else:
        df = pd.read_excel(filename, sheet_name=sheetname,header=None).iloc[:,:]
        flag=4
        col_start=0
        col_end=df.shape[1]
        row_start=0
        row_end=df.shape[0]

    
    df=df.replace(np.nan," ",regex=True)
    data=df.values
    print(filename[-1])
    if(filename[-1]=='s'):
        sheet = xlrd.open_workbook(filename, formatting_info=True).sheet_by_name(sheetname)
        merged_info=pd.DataFrame(sheet.merged_cells).values
        a=merged_info
        a=a[a[:,0].argsort()]
        

    elif(filename[-1]=='x'):
        wb = load_workbook(filename)
        sheet_ranges = wb[sheetname]
        a=sheet_ranges.merged_cells.ranges
        n=len(a)
        lst1=[]

        for i in range(0, n):
            lst2=[1,1,1,1]
            x = re.split(":", str(a[i]))    
            match = re.match(r"([a-z]+)([0-9]+)", x[0], re.I)
            lst2[0]= int(match.group(2))-1
            lst2[2]= _tonum(match.group(1))-1
    
            match = re.match(r"([a-z]+)([0-9]+)", x[1], re.I)
            lst2[1]= int(match.group(2))
            lst2[3]= _tonum(match.group(1))
    
            if((lst2[3]-lst2[2])==1 and (lst2[1]-lst2[0])==1):
                continue
    
            lst1.append(lst2)

        merged_info=np.array(lst1)
        a=merged_info
        a=a[a[:,0].argsort()]
        
    else:
        print("Check file type")
        return "error file type"
    
    a=a[np.where(a[:,0]<row_end)]
    a=a[np.where(a[:,0]>=row_start)]
    a=a[np.where(a[:,2]>=col_start)]
    a=a[np.where(a[:,2]<col_end)]

    print(a.tolist())
    a=np.subtract(a,[row_start,row_start,col_start,col_start])
    
    
    merged_info=a
    length1=merged_info.shape[0]
    

    dictionary={}
    for i in range(0,data.shape[0]):
        for j in range(0,data.shape[1]):
            a=[0,1,1]
            a[0]=data[i][j]
            dictionary.setdefault("row"+str(i), {})["col"+str(j)] =a

    for i in range(0,length1):
        x=merged_info[i][1]-merged_info[i][0]
        y=merged_info[i][3]-merged_info[i][2]
        dictionary["row"+str(merged_info[i][0])]["col"+str(merged_info[i][2])][1]=x
        dictionary["row"+str(merged_info[i][0])]["col"+str(merged_info[i][2])][2]=y

        for k in range(0,x):
            if(k!=0): 
                del dictionary["row"+str(merged_info[i][0]+k)]["col"+str(merged_info[i][2])]    
            for l in range(1,y):      
                del dictionary["row"+str(merged_info[i][0]+k)]["col"+str(merged_info[i][2]+l)]
    
    return dictionary


def toHtml(filename, sheetname,filesave,custom_row=False,custom_col=False,**kwargs):
    dictionary=toDictionary(filename,sheetname,custom_row=custom_row,custom_col=custom_col,**kwargs)
    print(dictionary)
    html="<html>\n<head>\n<script src=\"https://ajax.googleapis.com/ajax/libs/jquery/3.5.1/jquery.min.js\"></script>\n"
    html=html+"<style>\n div{\nalign:center;\n}\n"
    html=html+"table, td{\nborder-style:solid;\nborder-width: 2px;\nborder-collapse:collapse;\ntext-align:center;\n}\n"
    html=html+"</style>\n</head>"
    html=html+"<body>\n<h2>Table</h2>\n<form id=\"form_data\" action=\"\">\n<div id=\"tab\"></div>\n</form>\n</body>"

    html=html+"<script>\nvar myObj=[];\n{\nmyObj={\n"

    length2=len(dictionary)
    i=0
    text=""
    for k in dictionary:
        i=i+1
        if(i!=length2):
          html=html+"\n\""+k+"\":"+ str(dictionary[k])+" ,"
        else:
          html=html+"\n\""+k+"\":"+ str(dictionary[k])
    
    html=html+"};\n}"

    html=html+"\nvar txt=\"<table>\""+";\nvar ylen=0;\nfor(x in myObj)\n{\nylen++;\n}\n"
    html=html+"for(x in myObj)\n{\n"
    html=html+"txt+=\"<tr>\";\nfor(y in myObj[x])\n{\n "
    text=text+"\"<td class=\\\"\"+x+\" \"+y+\"\\\" id=\\\"\"+x+y+\"\\\""
    text=text+"rowspan=\\\"\"+myObj[x][y][1]+\"\\\" colspan=\\\"\"+myObj[x][y][2]+\"\\\">\"+myObj[x][y][0]+\"</td>\""
    
    html=html+"\ntxt+="+text

    html=html+";\n}\ntxt+=\"</tr>\";\n}\ntxt+=\"</table>\";\ndocument.getElementById(\"tab\").innerHTML=txt;\n</script>"

    
    f = open(filesave, "w",encoding="utf-8")
    f.write(html)
    f.close()
    print("DONE")


def _tonum(string):
    l=len(string)
    num=0
    for i in range(0,l):
        num=num + (ord(string[l-1-i])-64)*(26**i)
    return num




    


